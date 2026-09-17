"""
Servicio de Exportación a Excel (.xlsx) con Formato Profesional para Cuerpo Técnico.
Genera libros de Microsoft Excel estilizados con colores, bordes, alineación,
ajuste automático de ancho de columnas y congelación de cabeceras.
"""

import io
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd


def export_dataframe_to_formatted_excel(
    df: pd.DataFrame,
    sheet_name: str = "Planificación",
    header_color: str = "1E3A8A",
    freeze_header: bool = True
) -> bytes:
    """
    Convierte un DataFrame en un archivo binario .xlsx con formato visual idéntico
    al estándar de trabajo del preparador físico:
    - Cabecera con fondo azul marino (#1E3A8A), texto blanco negrita y centrado.
    - Columnas autoajustadas en anchura para evitar texto cortado.
    - Fila 'EQUIPO' destacada en azul suave con bordes reforzados.
    - Fila '⭐ JUGADOR TOP' destacada en tonos dorados/ámbar.
    - Cuadrícula de Excel visible y panel superior congelado.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    import re
    clean_title = re.sub(r'[\\/*?:\[\]]', '_', str(sheet_name)).strip()[:31]
    ws.title = clean_title if clean_title else "Hoja1"
    ws.views.sheetView[0].showGridLines = True

    if freeze_header:
        ws.freeze_panes = "A2"

    # Definición de bordes
    thin_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    thick_top_side = Side(border_style="medium", color="2563EB")
    double_bottom_side = Side(border_style="double", color="2563EB")
    team_border = Border(left=thin_side, right=thin_side, top=thick_top_side, bottom=double_bottom_side)

    # Estilos de cabecera
    clean_hex = header_color.replace("#", "")
    header_fill = PatternFill(start_color=clean_hex, end_color=clean_hex, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    # Estilos de filas especiales
    team_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    team_font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")

    top_player_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    top_player_font = Font(name="Calibri", size=10.5, bold=True, color="92400E")

    # Filas alternas normales
    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    normal_font = Font(name="Calibri", size=10.5, color="0F172A")
    bold_pos_font = Font(name="Calibri", size=10.5, bold=True, color="0F172A")

    # Excluir columnas técnicas auxiliares
    cols = [c for c in df.columns if not str(c).startswith("_")]

    # 1. Cabecera
    ws.row_dimensions[1].height = 26
    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(col_name))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = cell_border

    # 2. Filas de datos
    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        ws.row_dimensions[row_idx].height = 21
        pos_val = str(row.get("POSICIÓN", "")).strip().upper()
        row_type = str(row.get("_row_type", "")).strip().lower()

        is_team = "EQUIPO" in pos_val or row_type == "team"
        is_top = "⭐" in pos_val or "TOP" in pos_val or row_type == "top"

        for col_idx, col_name in enumerate(cols, start=1):
            val = row.get(col_name, "")

            # Formateo numérico inteligente
            cell_val = val
            num_fmt = None

            if isinstance(val, (int, float)):
                cell_val = val
                if isinstance(val, float):
                    num_fmt = "0.00" if ("km" in col_name.lower() or "velocidad" in col_name.lower()) else "0.0"
                else:
                    num_fmt = "#,##0"
            elif isinstance(val, str):
                s_val = val.strip()
                # Verificar si es número puro representado en texto
                clean_num = s_val.replace(",", ".")
                try:
                    if "." in clean_num:
                        cell_val = float(clean_num)
                        num_fmt = "0.00" if ("km" in col_name.lower() or "velocidad" in col_name.lower()) else "0.0"
                    elif clean_num.isdigit():
                        cell_val = int(clean_num)
                        num_fmt = "#,##0"
                except Exception:
                    cell_val = val

            cell = ws.cell(row=row_idx, column=col_idx, value=cell_val)
            if num_fmt:
                cell.number_format = num_fmt

            # Bordes
            cell.border = team_border if is_team else cell_border

            # Colores y tipografía
            if is_team:
                cell.fill = team_fill
                cell.font = team_font
            elif is_top:
                cell.fill = top_player_fill
                cell.font = top_player_font
            else:
                cell.fill = alt_row_fill if (row_idx % 2 == 0) else white_fill
                cell.font = bold_pos_font if col_name == "POSICIÓN" else normal_font

            # Alineaciones
            if col_name in ["POSICIÓN", "TIEMPO", "BASE EVALUACIÓN"]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_name == "JUGADOR":
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            elif isinstance(cell_val, (int, float)):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Ajuste de anchos de columna con margen extra para evitar truncados
    for col_idx, col_name in enumerate(cols, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_name))
        for r_idx in range(2, len(df) + 2):
            val_str = str(ws.cell(row=r_idx, column=col_idx).value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)

        # Reglas mínimas por tipo de columna conocida
        min_widths = {
            "POSICIÓN": 14,
            "JUGADOR": 26,
            "TIEMPO": 12,
            "BASE EVALUACIÓN": 18,
            "DISTANCIA TOTAL (km)": 22,
            "VELOCIDAD MAX (km/h)": 22,
            "HSR (m)": 14,
            "METROS EN SPRINT": 18,
            "#SPRINTS": 12,
            "#ACC EXPL": 13,
            "#DCC EXPL": 13,
        }
        calculated_width = max(max_len + 4, min_widths.get(col_name, 12))
        ws.column_dimensions[col_letter].width = calculated_width

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
